

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Tuple

from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[4]
PARAMETER_PATH = PROJECT / "数据" / "无人机应急物资运输基础数据" / "通信链路参数.xlsx"
EARTH_RADIUS_M = 6_371_000.0


@dataclass(frozen=True)
class DirectLinkParameters:
    frequency_mhz: float
    system_loss_db: float
    obstruction_loss_db: float
    sensitivity_dbm: float
    fade_margin_db: float
    uav_tx_dbm: float
    uav_gain_dbi: float
    gateway_tx_dbm: float
    gateway_gain_dbi: float
    gateway_height_m: float

    @property
    def effective_receive_threshold_dbm(self) -> float:
        return self.sensitivity_dbm + self.fade_margin_db

    @property
    def uplink_limit_db(self) -> float:
        return (
            self.uav_tx_dbm + self.uav_gain_dbi + self.gateway_gain_dbi
            - self.system_loss_db - self.effective_receive_threshold_dbm
        )

    @property
    def downlink_limit_db(self) -> float:
        return (
            self.gateway_tx_dbm + self.gateway_gain_dbi + self.uav_gain_dbi
            - self.system_loss_db - self.effective_receive_threshold_dbm
        )

    @property
    def bidirectional_limit_db(self) -> float:
        return min(self.uplink_limit_db, self.downlink_limit_db)


def load_direct_parameters(path: Path = PARAMETER_PATH) -> DirectLinkParameters:

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = {}
        for row in sheet.iter_rows(min_row=3, values_only=True):
            if row[0] and row[1] and row[4] is not None:
                values[(str(row[0]).strip(), str(row[1]).strip())] = float(row[4])

        def get(category: str, parameter: str) -> float:
            try:
                return values[(category, parameter)]
            except KeyError as exc:
                raise ValueError(f"通信参数表缺少 {category}/{parameter}") from exc

        return DirectLinkParameters(
            frequency_mhz=get("传播参数", "载波频率（MHz）"),
            system_loss_db=get("传播参数", "系统损耗（dB）"),
            obstruction_loss_db=get("传播参数", "地形遮挡附加损耗（dB）"),
            sensitivity_dbm=get("接收参数", "接收灵敏度（dBm）"),
            fade_margin_db=get("接收参数", "衰落裕量（dB）"),
            uav_tx_dbm=get("运输无人机", "发射功率（dBm）"),
            uav_gain_dbi=get("运输无人机", "天线增益（dBi）"),
            gateway_tx_dbm=get("固定网关 G01", "发射功率（dBm）"),
            gateway_gain_dbi=get("固定网关 G01", "天线增益（dBi）"),
            gateway_height_m=get("固定网关 G01", "天线离地高度（m）"),
        )
    finally:
        workbook.close()


def horizontal_distance_m(a: Tuple[float, float, float], b: Tuple[float, float, float]) -> float:

    lon1, lat1, _ = a
    lon2, lat2, _ = b
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = phi2 - phi1
    dlon = math.radians(lon2 - lon1)
    h = math.sin(dphi / 2) ** 2 + math.cos(phi1) * math.cos(phi2) * math.sin(dlon / 2) ** 2
    return 2 * EARTH_RADIUS_M * math.asin(min(1.0, math.sqrt(h)))


def distance_3d_m(a: Tuple[float, float, float], b: Tuple[float, float, float]) -> float:
    return math.hypot(horizontal_distance_m(a, b), b[2] - a[2])


def free_space_loss_db(distance_m: float, frequency_mhz: float) -> float:
    if distance_m < 0 or frequency_mhz <= 0:
        raise ValueError("三维距离不得为负且载波频率必须大于零")


    effective_distance_m = max(distance_m, 1.0)
    return 32.45 + 20 * math.log10(frequency_mhz) + 20 * math.log10(effective_distance_m / 1000.0)


def path_loss_db(fspl_db: float, blocked: bool, parameters: DirectLinkParameters) -> float:
    return fspl_db + (parameters.obstruction_loss_db if blocked else 0.0)
