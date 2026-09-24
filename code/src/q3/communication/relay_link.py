u"""运输无人机—中继、中继—G01 两类双向链路参数和预算。"""

from dataclasses import dataclass
from pathlib import Path
from typing import Tuple

from openpyxl import load_workbook

from .link_budget import (
    PARAMETER_PATH,
    distance_3d_m,
    free_space_loss_db,
    load_direct_parameters,
)
from .terrain_block import DemTerrain

RELAY_UAV_PATH = (
    Path(__file__).resolve().parents[4]
    / "数据" / "无人机应急物资运输基础数据" / "中继无人机数据.xlsx"
)


@dataclass(frozen=True)
class RelayLinkParameters:
    frequency_mhz: float
    system_loss_db: float
    obstruction_loss_db: float
    sensitivity_dbm: float
    fade_margin_db: float
    uav_tx_dbm: float
    uav_gain_dbi: float
    relay_access_tx_dbm: float
    relay_access_gain_dbi: float
    relay_backhaul_tx_dbm: float
    relay_backhaul_gain_dbi: float
    gateway_tx_dbm: float
    gateway_gain_dbi: float
    max_hover_agl_m: float
    hover_height_step_m: float = 20.0
    grid_pixel_step: int = 4

    @property
    def receive_threshold_dbm(self) -> float:
        return self.sensitivity_dbm + self.fade_margin_db

    @property
    def uav_relay_limit_db(self) -> float:
        a_to_b = (
            self.uav_tx_dbm + self.uav_gain_dbi + self.relay_access_gain_dbi
            - self.system_loss_db - self.receive_threshold_dbm
        )
        b_to_a = (
            self.relay_access_tx_dbm + self.relay_access_gain_dbi + self.uav_gain_dbi
            - self.system_loss_db - self.receive_threshold_dbm
        )
        return min(a_to_b, b_to_a)

    @property
    def relay_gateway_limit_db(self) -> float:
        a_to_b = (
            self.relay_backhaul_tx_dbm + self.relay_backhaul_gain_dbi + self.gateway_gain_dbi
            - self.system_loss_db - self.receive_threshold_dbm
        )
        b_to_a = (
            self.gateway_tx_dbm + self.gateway_gain_dbi + self.relay_backhaul_gain_dbi
            - self.system_loss_db - self.receive_threshold_dbm
        )
        return min(a_to_b, b_to_a)


def load_relay_link_parameters(
    path: Path = PARAMETER_PATH, relay_path: Path = RELAY_UAV_PATH
) -> RelayLinkParameters:
    """从通信参数表与中继无人机参数表分别读取收发设备数据。"""
    direct = load_direct_parameters(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = {}
        for row in workbook.active.iter_rows(min_row=3, values_only=True):
            if row[0] and row[1] and row[4] is not None:
                values[(str(row[0]).strip(), str(row[1]).strip())] = float(row[4])
    finally:
        workbook.close()

    def get(category: str, parameter: str) -> float:
        try:
            return values[(category, parameter)]
        except KeyError as exc:
            raise ValueError(f"通信参数表缺少 {category}/{parameter}") from exc

    relay_book = load_workbook(relay_path, read_only=True, data_only=True)
    try:
        relay_row = next(
            row for row in relay_book.active.iter_rows(min_row=3, values_only=True)
            if row[0] == "R"
        )
        max_agl = float(relay_row[18])
    finally:
        relay_book.close()

    return RelayLinkParameters(
        frequency_mhz=direct.frequency_mhz,
        system_loss_db=direct.system_loss_db,
        obstruction_loss_db=direct.obstruction_loss_db,
        sensitivity_dbm=direct.sensitivity_dbm,
        fade_margin_db=direct.fade_margin_db,
        uav_tx_dbm=get("运输无人机", "发射功率（dBm）"),
        uav_gain_dbi=get("运输无人机", "天线增益（dBi）"),
        relay_access_tx_dbm=get("中继接入端", "发射功率（dBm）"),
        relay_access_gain_dbi=get("中继接入端", "天线增益（dBi）"),
        relay_backhaul_tx_dbm=get("中继回传端", "发射功率（dBm）"),
        relay_backhaul_gain_dbi=get("中继回传端", "天线增益（dBi）"),
        gateway_tx_dbm=direct.gateway_tx_dbm,
        gateway_gain_dbi=direct.gateway_gain_dbi,
        max_hover_agl_m=max_agl,
    )


def evaluate_link(
    a: Tuple[float, float, float],
    b: Tuple[float, float, float],
    limit_db: float,
    parameters: RelayLinkParameters,
    terrain: DemTerrain,
) -> dict:
    terrain_result = terrain.check_line(a, b)
    fspl = free_space_loss_db(distance_3d_m(a, b), parameters.frequency_mhz)
    loss = fspl + (parameters.obstruction_loss_db if terrain_result.blocked else 0.0)
    return {
        "available": int(loss <= limit_db),
        "terrain_blocked": int(terrain_result.blocked),
        "fspl_db": fspl,
        "path_loss_db": loss,
        "margin_db": limit_db - loss,
    }
