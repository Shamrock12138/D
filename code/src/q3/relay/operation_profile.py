u"""Step7：中继飞行/能耗/时间 profile 与 gap job options 预计算（Pattern 级）。"""

import hashlib
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from src.dem_route import DEMRouteAnalyzer
from src.q2.battery import charge_time_to_full

PROJECT = Path(__file__).resolve().parents[3]
DATA = PROJECT / "data"

RELAY_UAV_XLSX = (
    Path(__file__).resolve().parents[4]
    / "数据" / "无人机应急物资运输基础数据" / "中继无人机数据.xlsx"
)
RELAY_BATTERY_CSV = DATA / "中继无人机_共享电池.csv"
SERVICE_AREA_CSV = DATA / "服务区数据.csv"

DEM_TIF = (
    Path(__file__).resolve().parents[4]
    / "数据" / "镇龙乡地理空间数据" / "镇龙乡及周边地理数据"
    / "数字高程模型数据（DEM）" / "镇龙乡及周边30米DEM.tif"
)

SITES_PATH = DATA / "q3_relay_sites.csv"
GAP_OPTIONS_PATH = DATA / "q3_gap_relay_options.csv"
GAPS_PATH = DATA / "q3_pattern_comm_gaps.csv"

OUT_PROFILES_PATH = DATA / "q3_relay_operation_profiles.csv"
OUT_JOB_OPTIONS_PATH = DATA / "q3_relay_job_options.csv"

G = 9.80665

OUT_MANIFEST_PATH = DATA / "q3_step7_manifest.json"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@dataclass(frozen=True)
class RelayFlightParams:
    cruise_speed_ms: float
    climb_speed_ms: float
    descend_speed_ms: float
    mass_kg: float
    cruise_power_kw: float
    climb_efficiency: float
    descend_efficiency: float
    energy_capacity_kwh: float
    safety_margin: float
    prep_time_s: float
    link_time_s: float
    turn_time_s: float
    hover_power_kw: float
    comm_power_kw: float
    full_charge_time_s: float

    @property
    def service_power_kw(self) -> float:
        return self.hover_power_kw + self.comm_power_kw

    @property
    def max_energy_kwh(self) -> float:
        return (1.0 - self.safety_margin) * self.energy_capacity_kwh


def load_relay_flight_parameters() -> RelayFlightParams:
    wb = load_workbook(RELAY_UAV_XLSX, read_only=True, data_only=True)
    try:
        ws = wb.active
        headers = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
        values = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        params = dict(zip(headers, values))
    finally:
        wb.close()

    battery = pd.read_csv(RELAY_BATTERY_CSV)
    rb_row = battery[battery["机型编号"] == "R"]
    if rb_row.empty:
        raise ValueError("中继无人机_共享电池.csv 中未找到机型 R")
    full_charge_time_s = float(rb_row["等效完全充电时间（s）"].values[0])

    return RelayFlightParams(
        cruise_speed_ms=float(params["计划巡航速度（m/s）"]),
        climb_speed_ms=float(params["最大爬升速度（m/s）"]),
        descend_speed_ms=float(params["最大下降速度（m/s）"]),
        mass_kg=float(params["计划起飞总质量（kg）"]),
        cruise_power_kw=float(params["巡航功率（kW）"]),
        climb_efficiency=float(params["爬升能耗效率"]),
        descend_efficiency=float(params["下降能耗效率"]),
        energy_capacity_kwh=float(params["能源组件可用能量（kWh）"]),
        safety_margin=float(params["返航电量下限（%）"]) / 100.0,
        prep_time_s=float(params["工位固定准备时间（s）"]),
        link_time_s=float(params["建链时间（s）"]),
        turn_time_s=float(params["架次周转时间（s）"]),
        hover_power_kw=float(params["悬停功率（kW）"]),
        comm_power_kw=float(params["通信附加功率（kW）"]),
        full_charge_time_s=full_charge_time_s,
    )


def load_o01() -> Tuple[float, float, float]:
    df = pd.read_csv(SERVICE_AREA_CSV)
    row = df[df["V"] == "O01"]
    if row.empty:
        raise ValueError("服务区数据中未找到 O01")
    r = row.iloc[0]
    return float(r["x"]), float(r["y"]), float(r["h"])


def build_site_operation_profiles(
    sites: pd.DataFrame,
    dem: DEMRouteAnalyzer,
    params: RelayFlightParams,
) -> pd.DataFrame:
    o01_lon, o01_lat, o01_ground = load_o01()
    o01_op = o01_ground
    safety_margin_m = 50.0

    rows = []
    failed_sites = []
    for idx, site in sites.iterrows():
        cid = site["candidate_id"]
        lon = float(site["lon"])
        lat = float(site["lat"])
        ground = float(site["ground_height"])
        agl = float(site["agl_height"])
        hover_abs = ground + agl
        distance_m = float(site["horizontal_distance_to_O01_m"])

        try:
            h_max_dem = dem.get_max_dem_along_route((o01_lon, o01_lat), (lon, lat))
        except RuntimeError as exc:
            failed_sites.append(cid)
            continue

        cruise_h = max(h_max_dem + safety_margin_m, o01_op, hover_abs)

        out_climb_m = cruise_h - o01_op
        out_descend_m = cruise_h - hover_abs
        ret_climb_m = cruise_h - hover_abs
        ret_descend_m = cruise_h - o01_op

        out_climb_s = out_climb_m / params.climb_speed_ms if params.climb_speed_ms > 0 else 0.0
        cruise_s = distance_m / params.cruise_speed_ms if params.cruise_speed_ms > 0 else 0.0
        out_descend_s = out_descend_m / params.descend_speed_ms if params.descend_speed_ms > 0 else 0.0

        out_time_s = out_climb_s + cruise_s + out_descend_s

        ret_climb_s = ret_climb_m / params.climb_speed_ms if params.climb_speed_ms > 0 else 0.0
        ret_descend_s = ret_descend_m / params.descend_speed_ms if params.descend_speed_ms > 0 else 0.0

        ret_time_s = ret_climb_s + cruise_s + ret_descend_s

        cruise_energy_kwh = params.cruise_power_kw * cruise_s / 3600.0

        if params.climb_efficiency > 0:
            out_climb_energy_kwh = (
                params.mass_kg * G * out_climb_m / (3.6e6 * params.climb_efficiency)
            )
            ret_climb_energy_kwh = (
                params.mass_kg * G * ret_climb_m / (3.6e6 * params.climb_efficiency)
            )
        else:
            out_climb_energy_kwh = 0.0
            ret_climb_energy_kwh = 0.0

        out_energy_kwh = out_climb_energy_kwh + cruise_energy_kwh
        ret_energy_kwh = ret_climb_energy_kwh + cruise_energy_kwh
        roundtrip_energy_kwh = out_energy_kwh + ret_energy_kwh

        lead_time_s = params.prep_time_s + out_time_s + params.link_time_s

        rows.append({
            "candidate_id": cid,
            "lon": lon,
            "lat": lat,
            "ground_height_m": ground,
            "agl_height_m": agl,
            "hover_absolute_height_m": hover_abs,
            "horizontal_distance_m": distance_m,
            "max_dem_height_m": h_max_dem,
            "cruise_height_m": cruise_h,
            "outbound_climb_m": out_climb_m,
            "outbound_descent_m": out_descend_m,
            "return_climb_m": ret_climb_m,
            "return_descent_m": ret_descend_m,
            "outbound_time_s": out_time_s,
            "return_time_s": ret_time_s,
            "outbound_energy_kWh": out_energy_kwh,
            "return_energy_kWh": ret_energy_kwh,
            "roundtrip_flight_energy_kWh": roundtrip_energy_kwh,
            "lead_time_s": lead_time_s,
        })

    if failed_sites:
        raise RuntimeError(
            f"DEM 栅格穿越失败，{len(failed_sites)} 个站点无法计算航线最高地形: "
            f"{failed_sites[:10]}{'...' if len(failed_sites) > 10 else ''}"
        )
    return pd.DataFrame(rows)


def build_gap_job_options(
    gap_options: pd.DataFrame,
    gaps: pd.DataFrame,
    sites: pd.DataFrame,
    profiles: pd.DataFrame,
    params: RelayFlightParams,
) -> pd.DataFrame:
    profile_map = profiles.set_index("candidate_id")

    gap_info = gaps[
        ["gap_id", "pattern_id", "tau_start", "tau_end", "coverage_start", "coverage_end"]
    ].copy()
    gap_info["service_duration_s"] = (
        gap_info["coverage_end"] - gap_info["coverage_start"]
    )

    options = gap_options[gap_options["full_cover"] == 1].copy()
    if len(options) != len(gap_options):
        removed = len(gap_options) - len(options)
        print(f"  移除非 full_cover option: {removed}")

    options = options.merge(gap_info, on="gap_id", how="left")
    options = options.merge(sites[["candidate_id", "lon", "lat", "ground_height", "agl_height"]],
                            on="candidate_id", how="left", suffixes=("", "_site"))

    rows = []
    for _, row in options.iterrows():
        cid = row["candidate_id"]
        if cid not in profile_map.index:
            continue
        prof = profile_map.loc[cid]

        cov_start = float(row["coverage_start"])
        cov_end = float(row["coverage_end"])
        gap_duration_s = float(row["service_duration_s"])
        service_energy_kwh = params.service_power_kw * (params.link_time_s + gap_duration_s) / 3600.0

        out_energy = float(prof["outbound_energy_kWh"])
        ret_energy = float(prof["return_energy_kWh"])
        total_energy_kwh = out_energy + service_energy_kwh + ret_energy

        if total_energy_kwh > params.max_energy_kwh:
            continue

        end_soc = 1.0 - total_energy_kwh / params.energy_capacity_kwh
        charge_time_s = charge_time_to_full(end_soc, params.full_charge_time_s)

        lead_time_s = float(prof["lead_time_s"])
        out_time_s = float(prof["outbound_time_s"])
        ret_time_s = float(prof["return_time_s"])

        dispatch_offset_s = cov_start - lead_time_s
        arrival_offset_s = cov_start - params.link_time_s
        service_start_offset_s = cov_start
        service_end_offset_s = cov_end
        return_offset_s = cov_end + ret_time_s

        min_transport_start_s = max(0.0, lead_time_s - cov_start)

        uav_occupancy_s = lead_time_s + gap_duration_s + ret_time_s + params.turn_time_s
        bat_occupancy_s = lead_time_s + gap_duration_s + ret_time_s + charge_time_s

        rows.append({
            "gap_id": row["gap_id"],
            "pattern_id": row["pattern_id"],
            "candidate_id": cid,
            "tau_start_s": float(row["tau_start"]),
            "tau_end_s": float(row["tau_end"]),
            "coverage_start_s": cov_start,
            "coverage_end_s": cov_end,
            "service_duration_s": gap_duration_s,
            "dispatch_offset_s": dispatch_offset_s,
            "arrival_offset_s": arrival_offset_s,
            "service_start_offset_s": service_start_offset_s,
            "service_end_offset_s": service_end_offset_s,
            "return_offset_s": return_offset_s,
            "min_transport_start_s": min_transport_start_s,
            "outbound_time_s": out_time_s,
            "return_time_s": ret_time_s,
            "lead_time_s": lead_time_s,
            "outbound_energy_kWh": out_energy,
            "return_energy_kWh": ret_energy,
            "service_energy_kWh": service_energy_kwh,
            "relay_energy_kWh": total_energy_kwh,
            "end_soc": end_soc,
            "charge_time_s": charge_time_s,
            "relay_uav_occupancy_s": uav_occupancy_s,
            "energy_component_occupancy_s": bat_occupancy_s,
            "lon": float(prof["lon"]),
            "lat": float(prof["lat"]),
            "agl_height_m": float(prof["agl_height_m"]),
            "min_access_margin_db": float(row["min_access_margin_db"]),
            "relay_g01_margin_db": float(row["relay_g01_margin_db"]),
        })

    return pd.DataFrame(rows)


def validate_step7(
    job_options: pd.DataFrame,
    sites: pd.DataFrame,
    profiles: pd.DataFrame,
    params: RelayFlightParams,
    all_gap_ids: set,
) -> bool:
    print("\n  Step7 完整性验证")
    print("  " + "-" * 50)
    all_ok = True

    covered_ids = set(job_options["gap_id"].unique())
    total_gaps = len(all_gap_ids)
    print(f"  Gaps with >=1 energy-feasible option: {len(covered_ids)} / {total_gaps}")
    missing = all_gap_ids - covered_ids
    if missing:
        sample = sorted(missing)[:20]
        print(f"  *** FAIL: {len(missing)} gaps 无能源可行 option: {sample}...")
        all_ok = False

    if (job_options["relay_energy_kWh"] > params.max_energy_kwh).any():
        violations = job_options[job_options["relay_energy_kWh"] > params.max_energy_kwh]
        print(f"  *** FAIL: {len(violations)} options 超过安全余量 {params.max_energy_kwh}kWh")
        all_ok = False
    else:
        print(f"  PASS: 所有 option relay_energy_kWh <= {params.max_energy_kwh}kWh")

    if (job_options["end_soc"] < params.safety_margin).any():
        violations = job_options[job_options["end_soc"] < params.safety_margin]
        print(f"  *** FAIL: {len(violations)} options end_soc < {params.safety_margin}")
        all_ok = False
    else:
        print(f"  PASS: 所有 option end_soc >= {params.safety_margin}")

    agl_ok = sites["agl_height"].between(0.1, 300.0).all()
    if agl_ok:
        print("  PASS: 所有 relay site 0 < AGL <= 300m")
    else:
        violations = sites[~sites["agl_height"].between(0.1, 300.0)]
        print(f"  *** FAIL: {len(violations)} sites AGL 越界")
        all_ok = False

    h_ok = (profiles["cruise_height_m"] >= profiles["max_dem_height_m"] + 50.0 - 1e-9).all()
    if h_ok:
        print("  PASS: 所有 cruise_height >= h_max_DEM + 50")
    else:
        print("  *** FAIL: cruise_height 不满足 h_max_DEM + 50")
        all_ok = False

    time_cols = [c for c in job_options.columns if "time_s" in c or "offset_s" in c or "occupancy_s" in c]
    for col in time_cols:
        if col in job_options.columns:
            col_data = pd.to_numeric(job_options[col], errors="coerce")
            if col_data.isna().any() or (col_data == np.inf).any() or (col_data == -np.inf).any():
                print(f"  *** FAIL: {col} 包含 NaN/inf")
                all_ok = False
    energy_cols = [c for c in job_options.columns if "energy" in c]
    for col in energy_cols:
        col_data = pd.to_numeric(job_options[col], errors="coerce")
        if (col_data < 0).any():
            print(f"  *** FAIL: {col} 包含负值")
            all_ok = False
        if col_data.isna().any() or (col_data == np.inf).any():
            print(f"  *** FAIL: {col} 包含 NaN/inf")
            all_ok = False

    if "charge_time_s" in job_options.columns:
        ct = pd.to_numeric(job_options["charge_time_s"], errors="coerce")
        if (ct < 0).any() or ct.isna().any():
            print("  *** FAIL: charge_time_s 包含负值/NaN")
            all_ok = False
        else:
            print("  PASS: charge_time_s >= 0 且无 NaN")

    if all_ok:
        print(f"\n  *** 全部验证 PASS ***")
    else:
        print(f"\n  *** 验证 FAIL ***")

    return all_ok


def run_step7():
    print("=" * 60)
    print("Step7: 中继飞行/能耗 profile 与 job options 预计算")
    print("=" * 60)

    print("\n[1/7] 加载参数...")
    params = load_relay_flight_parameters()
    o01_lon, o01_lat, o01_ground = load_o01()
    print(f"  O01: ({o01_lon}, {o01_lat}) 地面高 {o01_ground:.1f}m")
    print(f"  R 巡航 {params.cruise_speed_ms}m/s  爬升 {params.climb_speed_ms}m/s  下降 {params.descend_speed_ms}m/s")
    print(f"  R 电池 {params.energy_capacity_kwh}kWh  安全余量 {params.safety_margin*100:.0f}%")
    print(f"  R 最大任务能耗 {params.max_energy_kwh}kWh")

    print("\n[2/7] 加载输入数据...")
    sites = pd.read_csv(SITES_PATH)
    gap_options = pd.read_csv(GAP_OPTIONS_PATH)
    gaps = pd.read_csv(GAPS_PATH)
    print(f"  Relay sites:       {len(sites)}")
    print(f"  Gap alternatives:  {len(gap_options)}")
    print(f"  Gaps:              {len(gaps)}")

    full_cover_count = (gap_options["full_cover"] == 1).sum()
    assert full_cover_count == len(gap_options), \
        f"存在非 full_cover option: {len(gap_options) - full_cover_count}"
    print(f"  full_cover=1:      {full_cover_count} (全部)")

    print("\n[3/7] 初始化 DEM...")
    dem = DEMRouteAnalyzer()
    print(f"  DEM CRS: {dem.crs}, 分辨率: {dem.resolution}")

    print("\n[4/7] 对每个 relay site 计算 O01<->RP 飞行 profile (414 站点)...")
    profiles = build_site_operation_profiles(sites, dem, params)
    profiles.to_csv(OUT_PROFILES_PATH, index=False, encoding="utf-8-sig")
    print(f"  输出: {OUT_PROFILES_PATH} ({len(profiles)} rows)")

    dem.close()

    print("\n[5/7] 生成 gap job options...")
    job_options = build_gap_job_options(gap_options, gaps, sites, profiles, params)
    removed = len(gap_options) - len(job_options)
    print(f"  能源可行 option:     {len(job_options)}")
    print(f"  能源否决 (移除):     {removed}")

    print(f"\n  Relay job energy 统计:")
    stats = job_options["relay_energy_kWh"].describe()
    for key in ["count", "mean", "min", "25%", "50%", "75%", "max"]:
        if key in stats:
            print(f"    {key}: {stats[key]:.4f}")

    min_soc = job_options["end_soc"].min()
    print(f"\n  Min end SOC: {min_soc*100:.2f}%")

    job_options.to_csv(OUT_JOB_OPTIONS_PATH, index=False, encoding="utf-8-sig")
    print(f"\n  输出: {OUT_JOB_OPTIONS_PATH} ({len(job_options)} rows)")

    print("\n[6/7] 完整性验证...")
    all_gap_ids = set(gaps["gap_id"].unique())
    ok = validate_step7(job_options, sites, profiles, params, all_gap_ids)

    print("\n[7/7] 汇总")
    print(f"  Relay sites:                       {len(sites)}")
    print(f"  Input gap alternatives:            {len(gap_options)}")
    print(f"  Gaps:                              {len(all_gap_ids)}")
    print(f"  Energy-feasible alternatives:      {len(job_options)}")
    print(f"  Energy-infeasible removed:         {removed}")
    n_covered = job_options["gap_id"].nunique()
    lost = len(all_gap_ids) - n_covered
    print(f"  Gaps with >=1 feasible option:     {n_covered} / {len(all_gap_ids)}")
    print(f"  Gaps with 0 feasible option:       {lost}")
    print(f"  Max relay job energy:              {job_options['relay_energy_kWh'].max():.4f} kWh")
    print(f"  Min end SOC:                       {min_soc*100:.2f}%")

    manifest = {
        "step": "Step7",
        "description": "中继飞行/能耗/时间 profile 与 gap job options 预计算",
        "inputs": {
            "q3_relay_sites.csv": _sha256(SITES_PATH),
            "q3_gap_relay_options.csv": _sha256(GAP_OPTIONS_PATH),
            "q3_pattern_comm_gaps.csv": _sha256(GAPS_PATH),
            "中继无人机数据.xlsx": _sha256(RELAY_UAV_XLSX),
            "中继无人机_共享电池.csv": _sha256(RELAY_BATTERY_CSV),
            "服务区数据.csv": _sha256(SERVICE_AREA_CSV),
            "镇龙乡及周边30米DEM.tif": _sha256(DEM_TIF),
        },
        "outputs": {
            "q3_relay_operation_profiles.csv": _sha256(OUT_PROFILES_PATH),
            "q3_relay_job_options.csv": _sha256(OUT_JOB_OPTIONS_PATH),
        },
        "stats": {
            "relay_sites": len(sites),
            "gaps": len(all_gap_ids),
            "input_gap_alternatives": len(gap_options),
            "energy_feasible_alternatives": len(job_options),
            "energy_infeasible_removed": removed,
            "gaps_with_feasible_option": n_covered,
            "gaps_without_feasible_option": lost,
            "max_relay_job_energy_kWh": float(job_options["relay_energy_kWh"].max()),
            "min_end_soc": float(min_soc),
        },
    }
    with open(OUT_MANIFEST_PATH, "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, indent=2, ensure_ascii=False)
    print(f"\n  输出: {OUT_MANIFEST_PATH}")

    if not ok:
        print("\n*** 验证未通过，请检查输出后修复 ***")
        sys.exit(1)
    else:
        print("\n*** Step7 完成 ***")